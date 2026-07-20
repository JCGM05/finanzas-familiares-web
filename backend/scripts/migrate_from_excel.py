"""
Migra los datos de "Finanzas Familiares.xlsx" a la base de datos definida en app/models.py.

Uso:
    python scripts/migrate_from_excel.py "C:\\ruta\\a\\Finanzas Familiares.xlsx"

Lee Configuración, Movimientos, Objetivos y Patrimonio del Excel y crea las
filas correspondientes. Pensado para ejecutarse una sola vez sobre una base
de datos vacía (falla si ya hay categorías, para no duplicar).
"""

import sys
import warnings
from pathlib import Path

warnings.filterwarnings("ignore")

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

from app.database import SessionLocal, engine
from app.models import (
    AnnualExpense,
    Base,
    BalanceSnapshot,
    Category,
    Goal,
    GrupoCategoria,
    IncomeProfile,
    NetWorthEntry,
    NetWorthItem,
    PagadoCon,
    SuegrosSupportConfig,
    TipoCategoria,
    TipoMovimiento,
    TipoPatrimonio,
    Transaction,
    User,
)

TIPO_CATEGORIA_MAP = {
    "Recibo fijo": TipoCategoria.RECIBO_FIJO,
    "Variable": TipoCategoria.VARIABLE,
    "Anual": TipoCategoria.ANUAL,
    "Puntual": TipoCategoria.PUNTUAL,
    "Ingreso": TipoCategoria.INGRESO,
}

GRUPO_MAP = {
    "Necesidad": GrupoCategoria.NECESIDAD,
    "Deseo": GrupoCategoria.DESEO,
    "Puntual": GrupoCategoria.PUNTUAL,
    "Ingreso": GrupoCategoria.INGRESO,
}

TIPO_MOVIMIENTO_MAP = {
    "Gasto": TipoMovimiento.GASTO,
    "Ingreso Nómina": TipoMovimiento.INGRESO_NOMINA,
    "Ingreso extra": TipoMovimiento.INGRESO_EXTRA,
}

PAGADO_CON_MAP = {
    "Cuenta común": PagadoCon.CUENTA_COMUN,
    "Julio (personal)": PagadoCon.JULIO_PERSONAL,
    "Yanay (personal)": PagadoCon.YANAY_PERSONAL,
}

# Concepto de gasto anual -> nombre exacto de la categoría a la que provisiona.
ANNUAL_EXPENSE_CATEGORY_MAP = {
    "IBI (impuesto vivienda)": "IBI (impuesto vivienda)",
    "Seguro coche Julio": "Seguros de coche",
    "Seguro coche Yanay": "Seguros de coche",
}

MESES = [
    "Ene", "Feb", "Mar", "Abr", "May", "Jun",
    "Jul", "Ago", "Sep", "Oct", "Nov", "Dic",
]

PATRIMONIO_ANIO = 2026  # la hoja Patrimonio no indica año explícito; es el año en curso del Excel


def migrate(xlsx_path: str) -> None:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    cfg = wb["Configuración"]
    mov = wb["Movimientos"]
    obj = wb["Objetivos"]
    pat = wb["Patrimonio"]

    Base.metadata.create_all(bind=engine)
    db = SessionLocal()

    if db.query(Category).count() > 0:
        print("La base de datos ya tiene categorías: cancelo para no duplicar.")
        return

    # --- Usuarios ---
    julio = User(
        email="jcarlosgm05@gmail.com",
        password_hash="CAMBIAR_EN_CAPA_DE_LOGIN",
        display_name="Julio",
    )
    yanay = User(
        email="yanay@cambiar-este-email.local",
        password_hash="CAMBIAR_EN_CAPA_DE_LOGIN",
        display_name="Yanay",
    )
    db.add_all([julio, yanay])
    db.flush()

    # --- Ingresos (Configuración filas 8-9) ---
    db.add(
        IncomeProfile(
            user_id=julio.id,
            nomina_mensual=cfg["C8"].value,
            num_pagas=cfg["D8"].value,
            pct_a_comun=cfg["F8"].value,
        )
    )
    db.add(
        IncomeProfile(
            user_id=yanay.id,
            nomina_mensual=cfg["C9"].value,
            num_pagas=cfg["D9"].value,
            pct_a_comun=cfg["F9"].value,
            meses_pagas_extra=f"{cfg['E15'].value},{cfg['F15'].value}",
        )
    )

    # --- Saldo (Configuración filas 19-22) ---
    db.add(
        BalanceSnapshot(
            saldo_total=cfg["E19"].value,
            saldo_retenido=cfg["E20"].value or 0,
            nota_retencion="Retenido por compra de plaza de garaje pendiente de firmar",
            fecha=cfg["E22"].value.date(),
        )
    )

    # --- Categorías (Configuración filas 28-49) ---
    categorias_por_nombre: dict[str, Category] = {}
    for row_idx in range(28, 50):
        nombre = cfg[f"B{row_idx}"].value
        if not nombre:
            continue
        tipo_txt = cfg[f"C{row_idx}"].value
        grupo_txt = cfg[f"D{row_idx}"].value
        dia_cargo = cfg[f"E{row_idx}"].value
        presupuesto = cfg[f"F{row_idx}"].value or 0
        categoria = Category(
            nombre=nombre,
            tipo=TIPO_CATEGORIA_MAP[tipo_txt],
            grupo=GRUPO_MAP[grupo_txt],
            dia_cargo=dia_cargo,
            presupuesto_mensual=presupuesto,
            es_apoyo_suegros=nombre.startswith("Suegros ·"),
            orden=row_idx,
        )
        db.add(categoria)
        categorias_por_nombre[nombre] = categoria
    db.flush()

    # --- Gastos anuales (Configuración filas 55-57) ---
    for row_idx in range(55, 58):
        concepto = cfg[f"B{row_idx}"].value
        if not concepto:
            continue
        categoria_nombre = ANNUAL_EXPENSE_CATEGORY_MAP.get(concepto)
        db.add(
            AnnualExpense(
                category_id=categorias_por_nombre[categoria_nombre].id if categoria_nombre else None,
                concepto=concepto,
                importe_anual=cfg[f"C{row_idx}"].value,
                mes_cargo=cfg[f"D{row_idx}"].value,
                compania=cfg[f"F{row_idx}"].value,
            )
        )

    # --- Apoyo a los suegros (Configuración filas 62-63) ---
    db.add(
        SuegrosSupportConfig(
            meses_previstos=cfg["E62"].value,
            fecha_desde=cfg["E63"].value.date(),
        )
    )

    # --- Movimientos ---
    n_transacciones = 0
    for row in mov.iter_rows(min_row=7, max_row=mov.max_row):
        fecha = row[1].value  # columna B
        if fecha is None:
            continue
        tipo_txt = row[4].value  # E
        categoria_txt = row[5].value  # F
        descripcion = row[6].value  # G
        importe = row[7].value  # H
        pagado_con_txt = row[8].value  # I
        notas = row[9].value  # J

        categoria = categorias_por_nombre.get(categoria_txt)
        if categoria is None:
            print(f"  aviso: categoría desconocida '{categoria_txt}' en fila con fecha {fecha}, se omite")
            continue

        db.add(
            Transaction(
                fecha=fecha.date(),
                tipo=TIPO_MOVIMIENTO_MAP[tipo_txt],
                category_id=categoria.id,
                descripcion=descripcion,
                importe=importe,
                pagado_con=PAGADO_CON_MAP[pagado_con_txt],
                notas=notas,
                created_by_id=None,
            )
        )
        n_transacciones += 1

    # --- Objetivos (Objetivos: 6 tarjetas fijas -> filas) ---
    goal_cells = [
        ("B6", "B9", "B12", "AA33"),
        ("O6", "O9", "O12", "AA34"),
        ("AB6", "AB9", "AB12", "AA35"),
        ("B18", "B21", "B24", "AA36"),
        ("O18", "O21", "O24", "AA37"),
        ("AB18", "AB21", "AB24", "AA38"),
    ]
    for orden, (nombre_cell, meta_cell, ahorrado_cell, fecha_cell) in enumerate(goal_cells, start=1):
        nombre = obj[nombre_cell].value
        if not nombre:
            continue
        fecha_val = obj[fecha_cell].value
        db.add(
            Goal(
                nombre=nombre,
                meta=obj[meta_cell].value or 0,
                ahorrado=obj[ahorrado_cell].value or 0,
                fecha_objetivo=fecha_val.date() if fecha_val else None,
                orden=orden,
            )
        )

    # --- Patrimonio ---
    # Solo se controla la cuenta común (no cuentas personales de Julio/Yanay).
    activos = [
        ("Cuenta común", 8),
        ("Ahorro / Depósitos", 11), ("Inversiones", 12), ("Valor de la vivienda", 13),
        ("Otros activos", 14),
    ]
    pasivos = [
        ("Hipoteca pendiente", 18), ("Otros préstamos", 19), ("Tarjetas de crédito", 20),
    ]
    col_letters = ["C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N"]

    for orden, (nombre, row_idx) in enumerate(activos, start=1):
        item = NetWorthItem(nombre=nombre, tipo=TipoPatrimonio.ACTIVO, orden=orden)
        db.add(item)
        db.flush()
        for mes, col in enumerate(col_letters, start=1):
            db.add(NetWorthEntry(item_id=item.id, anio=PATRIMONIO_ANIO, mes=mes, importe=pat[f"{col}{row_idx}"].value or 0))

    for orden, (nombre, row_idx) in enumerate(pasivos, start=1):
        item = NetWorthItem(nombre=nombre, tipo=TipoPatrimonio.PASIVO, orden=orden)
        db.add(item)
        db.flush()
        for mes, col in enumerate(col_letters, start=1):
            db.add(NetWorthEntry(item_id=item.id, anio=PATRIMONIO_ANIO, mes=mes, importe=pat[f"{col}{row_idx}"].value or 0))

    db.commit()
    db.close()

    print(f"Migración completa: {len(categorias_por_nombre)} categorías, {n_transacciones} movimientos.")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Uso: python migrate_from_excel.py <ruta al xlsx>")
        sys.exit(1)
    migrate(sys.argv[1])
